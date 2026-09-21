"""Geometry QA: overflow past the safe area, and picture/card overlaps."""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu
IN = lambda v: Emu(v).inches
SAFE_B, SW = 6.90, 13.33

def boxes(s):
    for sh in s.shapes:
        yield sh, IN(sh.left), IN(sh.top), IN(sh.left+sh.width), IN(sh.top+sh.height)

def overlap(a, b):
    return not (a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1])

# default: the deck build_deck.py writes; or pass any .pptx:  python qa_geom.py path/to/deck.pptx
DECK = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "build" / "FraudLens_SIH2026_PixelRex.pptx"
prs = Presentation(str(DECK))
bad = 0
for i, s in enumerate(prs.slides, 1):
    msgs = []
    items = list(boxes(s))
    for sh, l, t, r, b in items:
        if sh.name in ("Shape 0", "Text 1", "Text 2", "Image 0"):
            continue
        if b > SAFE_B + 0.01:
            msgs.append(f"overflow: {sh.name} bottom={b:.2f}")
        if r > SW - 0.15 or l < 0.30:
            msgs.append(f"edge: {sh.name} l={l:.2f} r={r:.2f}")
    # pictures must not sit under a filled card
    pics = [(sh, l, t, r, b) for sh, l, t, r, b in items if sh.shape_type == 13
            and sh.name != "Image 0"]
    cards = [(sh, l, t, r, b) for sh, l, t, r, b in items
             if sh.shape_type == 1 and sh.name.startswith("Rounded")]
    for ps, *pb in pics:
        for cs, *cb in cards:
            if overlap(pb, cb):
                msgs.append(f"picture/card overlap: {ps.name} x {cs.name}")
    print(f"slide {i}: " + ("OK" if not msgs else ""))
    for m in msgs:
        print("   ", m); bad += 1
print("\nissues:", bad)
